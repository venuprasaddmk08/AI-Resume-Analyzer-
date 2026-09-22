from io import BytesIO

import pymupdf
import pytest
from docx import Document

from services.resume_parser import parse_resume


def _build_pdf_bytes(lines_with_sizes: list[tuple[str, float]]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 50
    for text, size in lines_with_sizes:
        page.insert_text((50, y), text, fontsize=size)
        y += size + 10
    data = doc.tobytes()
    doc.close()
    return data


def _build_blank_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def _build_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Jane Doe", level=1)
    document.add_paragraph("Experience", style="Heading 2")
    document.add_paragraph("Backend Developer at Acme Corp")
    document.add_paragraph("Education", style="Heading 2")
    document.add_paragraph("B.Tech in Computer Science")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Level"
    table.rows[1].cells[0].text = "Python"
    table.rows[1].cells[1].text = "Advanced"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_pdf_extracts_text_and_sections():
    content = _build_pdf_bytes(
        [
            ("Jane Doe", 18),
            ("Experience", 14),
            ("Backend Developer at Acme Corp using Python and FastAPI", 10),
            ("Education", 14),
            ("B.Tech in Computer Science", 10),
        ]
    )

    result = parse_resume(content, "resume.pdf")

    assert result.file_type == "pdf"
    assert result.page_count == 1
    assert "Backend Developer at Acme Corp" in result.normalized_text
    assert result.warnings == []
    section_names = {s.section for s in result.sections_detected}
    assert "experience" in section_names
    assert "education" in section_names


def test_parse_pdf_with_no_extractable_text_warns():
    content = _build_blank_pdf_bytes()

    result = parse_resume(content, "blank.pdf")

    assert result.normalized_text == ""
    assert "Text could not be extracted from this PDF." in result.warnings


def test_parse_pdf_rejects_malformed_content():
    with pytest.raises(ValueError):
        parse_resume(b"this is not a real pdf", "fake.pdf")


def test_parse_docx_extracts_paragraphs_headings_and_tables():
    content = _build_docx_bytes()

    result = parse_resume(content, "resume.docx")

    assert result.file_type == "docx"
    assert "Backend Developer at Acme Corp" in result.normalized_text
    assert "Python" in result.normalized_text
    section_names = {s.section for s in result.sections_detected}
    assert "experience" in section_names
    assert "education" in section_names
    assert len(result.tables) == 1
    assert result.tables[0][1] == ["Python", "Advanced"]
    assert any("table" in w.lower() for w in result.warnings)


def test_parse_docx_rejects_malformed_content():
    with pytest.raises(ValueError):
        parse_resume(b"not a real docx", "fake.docx")


def test_parse_txt_extracts_text_and_sections():
    content = b"Jane Doe\nSkills\nPython, SQL, FastAPI\nEducation\nB.Tech in CS"

    result = parse_resume(content, "resume.txt")

    assert result.file_type == "txt"
    assert "Python, SQL, FastAPI" in result.normalized_text
    section_names = {s.section for s in result.sections_detected}
    assert "skills" in section_names
    assert "education" in section_names


def test_parse_txt_handles_invalid_utf8_gracefully():
    content = b"Jane Doe\xff\xfeSkills: Python"

    result = parse_resume(content, "resume.txt")

    assert result.file_type == "txt"
    assert result.normalized_text != ""
    assert any("could not be decoded" in w for w in result.warnings)


def test_parse_txt_empty_file_warns():
    result = parse_resume(b"   \n\n  ", "empty.txt")

    assert result.normalized_text == ""
    assert "No text could be extracted from this document." in result.warnings


def test_parse_resume_rejects_unsupported_extension():
    with pytest.raises(ValueError):
        parse_resume(b"hello", "resume.exe")
