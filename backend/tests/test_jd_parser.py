from io import BytesIO

import pymupdf
import pytest
from docx import Document

from services.jd_parser import parse_job_description


def _build_pdf_bytes(lines_with_sizes: list[tuple[str, float]]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 50
    for text, size in lines_with_sizes:
        page.insert_text((50, y), text, fontsize=size)
        y += size + 10
    data = doc.tobytes()
    doc.close()
    return data


def _build_blank_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def _build_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Backend Developer", level=1)
    document.add_paragraph("Responsibilities", style="Heading 2")
    document.add_paragraph("Build and maintain REST APIs using Python and FastAPI.")
    document.add_paragraph("Requirements", style="Heading 2")
    document.add_paragraph("2+ years of experience with Python and SQL.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Level"
    table.rows[1].cells[0].text = "Docker"
    table.rows[1].cells[1].text = "Preferred"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_jd_pdf_extracts_text_and_sections():
    content = _build_pdf_bytes(
        [
            ("Backend Developer", 18),
            ("Responsibilities", 14),
            ("Build and maintain REST APIs using Python and FastAPI", 10),
            ("Requirements", 14),
            ("2+ years of experience with Python and SQL", 10),
        ]
    )

    result = parse_job_description(content=content, filename="jd.pdf")

    assert result.source == "pdf"
    assert result.page_count == 1
    assert "Build and maintain REST APIs" in result.normalized_text
    assert result.warnings == []
    section_names = {s.section for s in result.sections_detected}
    assert "responsibilities" in section_names
    assert "requirements" in section_names


def test_parse_jd_pdf_with_no_extractable_text_warns():
    content = _build_blank_pdf_bytes()

    result = parse_job_description(content=content, filename="blank.pdf")

    assert result.normalized_text == ""
    assert "Text could not be extracted from this PDF." in result.warnings


def test_parse_jd_pdf_rejects_malformed_content():
    with pytest.raises(ValueError):
        parse_job_description(content=b"not a real pdf", filename="fake.pdf")


def test_parse_jd_docx_extracts_paragraphs_headings_and_tables():
    content = _build_docx_bytes()

    result = parse_job_description(content=content, filename="jd.docx")

    assert result.source == "docx"
    assert "Python and FastAPI" in result.normalized_text
    section_names = {s.section for s in result.sections_detected}
    assert "responsibilities" in section_names
    assert "requirements" in section_names
    assert len(result.tables) == 1
    assert result.tables[0][1] == ["Docker", "Preferred"]
    assert any("table" in w.lower() for w in result.warnings)


def test_parse_jd_docx_rejects_malformed_content():
    with pytest.raises(ValueError):
        parse_job_description(content=b"not a real docx", filename="fake.docx")


def test_parse_jd_txt_extracts_text_and_sections():
    content = b"Backend Developer\nRequirements\nPython, SQL, FastAPI\nBenefits\nRemote work"

    result = parse_job_description(content=content, filename="jd.txt")

    assert result.source == "txt"
    assert "Python, SQL, FastAPI" in result.normalized_text
    section_names = {s.section for s in result.sections_detected}
    assert "requirements" in section_names
    assert "benefits" in section_names


def test_parse_jd_pasted_text_extracts_sections():
    text = "Backend Developer\nResponsibilities\nDesign REST APIs\nRequirements\nPython and SQL required"

    result = parse_job_description(pasted_text=text)

    assert result.source == "pasted"
    assert "Design REST APIs" in result.normalized_text
    section_names = {s.section for s in result.sections_detected}
    assert "responsibilities" in section_names
    assert "requirements" in section_names


def test_parse_jd_pasted_empty_text_warns():
    result = parse_job_description(pasted_text="   \n\n  ")

    assert result.normalized_text == ""
    assert "No text was provided." in result.warnings


def test_parse_jd_rejects_unsupported_extension():
    with pytest.raises(ValueError):
        parse_job_description(content=b"hello", filename="jd.exe")


def test_parse_jd_requires_file_or_pasted_text():
    with pytest.raises(ValueError):
        parse_job_description()
