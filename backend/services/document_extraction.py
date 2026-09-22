"""Generic, domain-agnostic document layout extraction.

Shared by resume_parser.py and jd_parser.py so the PyMuPDF/python-docx
extraction code exists in exactly one place. This module knows nothing
about resumes or job descriptions specifically — it only extracts text,
positions, and font sizes. Domain-specific section-heading vocabularies
live in the callers.
"""

from io import BytesIO
from typing import Optional

import pymupdf
from docx import Document

from schemas import Block


def median(values: list[float]) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    n = len(ordered)
    mid = n // 2
    if n % 2:
        return ordered[mid]
    return (ordered[mid - 1] + ordered[mid]) / 2


def match_section(line_text: str, keyword_map: dict[str, list[str]]) -> Optional[str]:
    """Exact-match (not substring) heading lookup so ordinary sentences that
    happen to start with a keyword aren't misread as section headings."""
    normalized = line_text.strip().lower().rstrip(":").strip()
    if not normalized or len(normalized) > 40:
        return None
    for canonical, phrases in keyword_map.items():
        if normalized in phrases:
            return canonical
    return None


def flatten_tables(tables: list[list[list[str]]]) -> list[str]:
    lines: list[str] = []
    for table in tables:
        for row in table:
            row_text = " | ".join(cell for cell in row if cell)
            if row_text:
                lines.append(row_text)
    return lines


def extract_pdf(content: bytes) -> tuple[list[Block], int, list[str]]:
    """Returns (blocks, page_count, warnings). Raises ValueError if not a valid PDF."""
    if not content.startswith(b"%PDF"):
        raise ValueError("File does not appear to be a valid PDF.")

    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("File does not appear to be a valid PDF.") from exc

    blocks: list[Block] = []
    font_sizes: list[float] = []

    try:
        page_count = doc.page_count
        for page_index in range(page_count):
            page = doc[page_index]
            page_dict = page.get_text("dict")
            for raw_block in page_dict.get("blocks", []):
                if raw_block.get("type") != 0:  # 0 = text block, skip images
                    continue
                for line in raw_block.get("lines", []):
                    spans = line.get("spans", [])
                    if not spans:
                        continue
                    line_text = "".join(span.get("text", "") for span in spans).strip()
                    if not line_text:
                        continue
                    max_font = max((span.get("size", 0) for span in spans), default=0.0)
                    bbox = [round(v, 2) for v in line.get("bbox", [0, 0, 0, 0])]
                    font_sizes.append(max_font)
                    blocks.append(
                        Block(
                            text=line_text,
                            page_number=page_index + 1,
                            bbox=bbox,
                            font_size=round(max_font, 1),
                            is_heading=False,
                        )
                    )
    finally:
        doc.close()

    body_font = median(font_sizes)
    for block in blocks:
        if block.font_size and body_font and block.font_size >= body_font + 1.5 and len(block.text) <= 60:
            block.is_heading = True

    return blocks, page_count, []


def extract_docx(content: bytes) -> tuple[list[Block], list[list[list[str]]], list[str]]:
    """Returns (blocks, tables, warnings). Raises ValueError if not a valid DOCX."""
    if not content.startswith(b"PK"):
        raise ValueError("File does not appear to be a valid DOCX document.")

    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("File does not appear to be a valid DOCX document.") from exc

    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        style_says_heading = style_name.lower().startswith("heading") or style_name.lower() == "title"
        blocks.append(Block(text=text, page_number=None, bbox=None, font_size=None, is_heading=style_says_heading))

    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    warnings: list[str] = []
    if tables:
        warnings.append(
            f"Document contains {len(tables)} table(s); table reading order is not guaranteed to match visual layout."
        )

    return blocks, tables, warnings


def extract_txt(content: bytes) -> tuple[list[Block], list[str]]:
    """Returns (blocks, warnings)."""
    warnings: list[str] = []
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("utf-8", errors="replace")
        warnings.append("Some characters in this file could not be decoded as UTF-8 and were replaced.")

    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    blocks = [Block(text=line, page_number=None, bbox=None, font_size=None, is_heading=False) for line in lines]
    return blocks, warnings
