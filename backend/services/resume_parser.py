"""Deterministic resume text/layout extraction for PDF, DOCX, and TXT.

No AI involved here — this only extracts what is literally present in the
uploaded file (text, positions, font sizes where available, and a heuristic
guess at section headings). Downstream AI analysis (Phase 4) builds on top
of this, not the other way around.
"""

from io import BytesIO
from typing import Optional

import pymupdf
from docx import Document

from schemas import Block, ParsedResume, SectionPosition

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}

# Canonical section name -> exact heading phrases (lowercase, no trailing punctuation)
# that count as a match. Kept as exact matches (not substring) so ordinary
# sentences that happen to start with a keyword aren't misread as headings.
SECTION_KEYWORDS: dict[str, list[str]] = {
    "summary": ["summary", "profile", "objective", "about me", "professional summary"],
    "experience": [
        "experience",
        "work experience",
        "employment history",
        "professional experience",
    ],
    "internships": ["internship", "internships"],
    "education": ["education", "academic background", "educational qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "key skills"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses & certifications"],
    "achievements": ["achievements", "awards", "honors", "honours"],
    "languages": ["languages"],
    "publications": ["publications"],
}


def _match_section(line_text: str) -> Optional[str]:
    normalized = line_text.strip().lower().rstrip(":").strip()
    if not normalized or len(normalized) > 40:
        return None
    for canonical, phrases in SECTION_KEYWORDS.items():
        if normalized in phrases:
            return canonical
    return None


def _median(values: list[float]) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    n = len(ordered)
    mid = n // 2
    if n % 2:
        return ordered[mid]
    return (ordered[mid - 1] + ordered[mid]) / 2


def parse_resume(content: bytes, filename: str) -> ParsedResume:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: .{extension or 'unknown'}")

    if extension == "pdf":
        return _parse_pdf(content)
    if extension == "docx":
        return _parse_docx(content)
    return _parse_txt(content)


def _parse_pdf(content: bytes) -> ParsedResume:
    if not content.startswith(b"%PDF"):
        raise ValueError("File does not appear to be a valid PDF.")

    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # pymupdf raises its own error types
        raise ValueError("File does not appear to be a valid PDF.") from exc

    blocks: list[Block] = []
    text_lines: list[str] = []
    font_sizes: list[float] = []

    try:
        page_count = doc.page_count
        for page_index in range(page_count):
            page = doc[page_index]
            page_dict = page.get_text("dict")
            for raw_block in page_dict.get("blocks", []):
                if raw_block.get("type") != 0:  # 0 = text block, skip images
                    continue
                for line in raw_block.get("lines", []):
                    spans = line.get("spans", [])
                    if not spans:
                        continue
                    line_text = "".join(span.get("text", "") for span in spans).strip()
                    if not line_text:
                        continue
                    max_font = max((span.get("size", 0) for span in spans), default=0.0)
                    bbox = [round(v, 2) for v in line.get("bbox", [0, 0, 0, 0])]
                    font_sizes.append(max_font)
                    text_lines.append(line_text)
                    blocks.append(
                        Block(
                            text=line_text,
                            page_number=page_index + 1,
                            bbox=bbox,
                            font_size=round(max_font, 1),
                            is_heading=False,
                        )
                    )
    finally:
        doc.close()

    body_font = _median(font_sizes)
    sections: list[SectionPosition] = []
    for block in blocks:
        if block.font_size and body_font and block.font_size >= body_font + 1.5 and len(block.text) <= 60:
            block.is_heading = True
        matched = _match_section(block.text)
        if matched:
            block.is_heading = True
            sections.append(SectionPosition(section=matched, page_number=block.page_number, text=block.text))

    normalized_text = "\n".join(text_lines)
    warnings: list[str] = []
    if not normalized_text.strip():
        warnings.append("Text could not be extracted from this PDF.")

    return ParsedResume(
        file_type="pdf",
        normalized_text=normalized_text,
        page_count=page_count,
        blocks=blocks,
        tables=[],
        sections_detected=sections,
        warnings=warnings,
    )


def _parse_docx(content: bytes) -> ParsedResume:
    if not content.startswith(b"PK"):
        raise ValueError("File does not appear to be a valid DOCX document.")

    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("File does not appear to be a valid DOCX document.") from exc

    blocks: list[Block] = []
    sections: list[SectionPosition] = []
    text_lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        style_says_heading = style_name.lower().startswith("heading") or style_name.lower() == "title"
        matched = _match_section(text)
        is_heading = style_says_heading or matched is not None

        text_lines.append(text)
        blocks.append(Block(text=text, page_number=None, bbox=None, font_size=None, is_heading=is_heading))
        if matched:
            sections.append(SectionPosition(section=matched, page_number=None, text=text))

    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                text_lines.append(row_text)
        tables.append(rows)

    normalized_text = "\n".join(text_lines)
    warnings: list[str] = []
    if not normalized_text.strip():
        warnings.append("No text could be extracted from this document.")
    if tables:
        warnings.append(
            f"Document contains {len(tables)} table(s); table reading order is not guaranteed to match visual layout."
        )

    return ParsedResume(
        file_type="docx",
        normalized_text=normalized_text,
        page_count=None,
        blocks=blocks,
        tables=tables,
        sections_detected=sections,
        warnings=warnings,
    )


def _parse_txt(content: bytes) -> ParsedResume:
    warnings: list[str] = []
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("utf-8", errors="replace")
        warnings.append("Some characters in this file could not be decoded as UTF-8 and were replaced.")

    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]

    blocks: list[Block] = []
    sections: list[SectionPosition] = []
    for line in lines:
        matched = _match_section(line)
        blocks.append(Block(text=line, page_number=None, bbox=None, font_size=None, is_heading=matched is not None))
        if matched:
            sections.append(SectionPosition(section=matched, page_number=None, text=line))

    normalized_text = "\n".join(lines)
    if not normalized_text.strip():
        warnings.append("No text could be extracted from this document.")

    return ParsedResume(
        file_type="txt",
        normalized_text=normalized_text,
        page_count=None,
        blocks=blocks,
        tables=[],
        sections_detected=sections,
        warnings=warnings,
    )
